# -*- coding: utf-8 -*-
"""T-05 报告生成器：REQ-002 合并版（简要+详细）组装（tasks.md v1.0）。

依赖链：factor_engine（T-04）+ trade_cal（E-7）+ interpretation（T-06）+ compliance（T-01）。
产出：HTML 在线版（图片 base64 自包含）+ docx 文件版，双形态同源。

DoD：
1. 章节完整性：REQ-002 清单一一对应（AC-001，check_sections 机械检查）；
2. 分档着色 BR-007（AC-004，band 三色）；
3. 复核日 = 出分日 + 5 交易日（AC-006，含跨节假日容错）；
4. "怎么读这个分数""风险与边界"章节在位（AC-010）；
5. 组装最后一步强制 T-01 合规扫描，FAIL 即阻断输出（AC-005）。
"""

from __future__ import annotations

import base64
import io
import os
from datetime import datetime as _dt
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd

from src.compliance import ComplianceBlockError, assert_compliant_or_raise, scan
from src.factor_engine import (FACTOR_TABLE, SIDE_ORDER, band, compute_H,
                               compute_pool_scores, side_means)
from src.interpretation import DiagnosisContext, get_interpretation
from src.trade_cal import early_review_flags, review_date

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
plt.rcParams["font.family"] = "Microsoft YaHei"
plt.rcParams["axes.unicode_minus"] = False

from docx import Document                     # noqa: E402
from docx.shared import Pt, RGBColor          # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402

# ---- REQ-002 章节标题（AC-001 机械检查清单） -----------------------------------
SECTION_TITLES: Dict[str, str] = {
    "meta":      "诊断对象与数据口径",
    "summary":   "诊断汇总",
    "oneline":   "第一层 · 一句话结论",
    "sides":     "四个观察侧面",
    "curve":     "股性曲线与历史位置 H",
    "howto":     "怎么读这个分数",
    "plain":     "第二层 · 白话解释",
    "boundary":  "风险与边界",
    "objective": "因子视角的客观评估",
    "compare":   "当前状态的两种比较",
    "dist":      "分项信号的分布",
    "dims":      "宽口径维度观察",
    "histpos":   "自身历史位置与持续性",
    "scorepos":  "当前得分在过去出现的位置",
    "stat":      "统计证据与解释边界",
    "factors":   "匿名因子明细",
    "math":      "数学方法与数据覆盖",
}

BOUNDARY_TEXT = (
    "本报告为量化因子状态的客观描述工具。分数描述因子状态，不代表上涨概率、可能性或收益预期，"
    "亦不构成任何证券交易类操作建议或投资顾问服务。四个观察侧面仅用于拆解展示，不另设权重；"
    "风险维度得分高不等于绝对风险低。数据来自公开渠道，可能存在延迟或修正，请以交易所披露为准。"
)
HOWTO_TEXT = ("分数为同池相对强弱的百分位打分（0–100），仅代表当前因子状态的相对位置，"
              "不代表概率、可能性或收益预期。高于 60 读取为偏有利信号，低于 40 读取为偏不利信号，"
              "40–60 为中性区。综合分 S 是 18 项因子等权合成后放回同池重新排名；H 是此前 52 个完成周 "
              "S 值的相对位置，两者口径不同，不可直接相加或平均。")


def _fmt(x, nd=0):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "无数据"
    return f"{x:.{nd}f}"


def _pct(x):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "无数据"
    return f"{x:+.1%}"


# ---- 图表（BR-007 三色；matplotlib 色值带 # 前缀） -------------------------------

def _png_b64(fig) -> str:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=130, bbox_inches="tight")
    plt.close(fig)
    return base64.b64encode(buf.getvalue()).decode()


def _chart_price(panel: pd.DataFrame, name: str) -> str:
    c = panel["close"]
    fig, ax = plt.subplots(figsize=(7.4, 2.8))
    ax.plot(c.index, c.values, color="#333333", lw=1.0, label="收盘价（前复权）")
    ax.plot(c.index, c.rolling(20).mean().values, color="#888888", lw=0.8, ls="--", label="MA20")
    ax.plot(c.index, c.rolling(60).mean().values, color="#aaaaaa", lw=0.8, ls=":", label="MA60")
    ax.set_title(f"{name} 股性曲线（日线）", fontsize=10)
    ax.legend(fontsize=8, loc="best")
    ax.grid(alpha=0.25)
    return _png_b64(fig)


def _chart_factors(fscores: pd.Series, labels: Dict[str, str]) -> str:
    fig, ax = plt.subplots(figsize=(7.4, 4.6))
    fids = [f[0] for f in FACTOR_TABLE]
    vals = [fscores.get(f, np.nan) for f in fids]
    names = [f"{f}·{labels.get(f, f)}" for f in fids]
    colors = [band(v)[1] for v in vals]
    ax.barh(range(len(fids)), [v if v == v else 0 for v in vals], color=colors)
    ax.set_yticks(range(len(fids)))
    ax.set_yticklabels(names, fontsize=7.5)
    ax.axvline(60, color="#C00000", lw=0.7, ls="--")
    ax.axvline(40, color="#1E7B3C", lw=0.7, ls="--")
    ax.set_xlim(0, 100)
    ax.set_title("18 项观察因子同池百分位（红=偏有利 >60 / 绿=偏不利 <40 / 灰=中性）", fontsize=9)
    return _png_b64(fig)


def _chart_s_series(series: pd.Series, name: str) -> str:
    fig, ax = plt.subplots(figsize=(7.4, 2.8))
    ax.plot(series.index, series.values, color="#555555", lw=1.2, marker="o", ms=2.5)
    ax.axhline(60, color="#C00000", lw=0.7, ls="--")
    ax.axhline(40, color="#1E7B3C", lw=0.7, ls="--")
    ax.set_ylim(0, 105)
    ax.set_title(f"{name} 综合分 S 周频轨迹（52 完成周，正式口径回算）", fontsize=10)
    ax.grid(alpha=0.25)
    return _png_b64(fig)


def _chart_hist(series: pd.Series, S_now: float) -> str:
    hist = series.iloc[:-1]
    fig, ax = plt.subplots(figsize=(7.4, 2.8))
    ax.hist(hist.dropna().values, bins=20, color="#999999")
    ax.axvline(S_now, color="#C00000", lw=1.4, label=f"当前 S={S_now:.0f}")
    ax.set_title("当前 S 在此前 52 个完成周 S 分布中的位置（H）", fontsize=10)
    ax.legend(fontsize=8)
    return _png_b64(fig)


# ---- 主流程 ----------------------------------------------------------------------

def build_report(panels: Dict[str, pd.DataFrame], target_code: str, target_name: str,
                 pool_name: str, out_dir: str,
                 weeks: int = 52) -> Dict:
    """组装合并版报告。返回 result dict；任何合规 FAIL 抛 ComplianceBlockError（不落盘）。"""
    os.makedirs(out_dir, exist_ok=True)
    target = panels[target_code]
    observe_date = target.index[-1]
    score_date = observe_date.date()

    # 1) 因子与分数（T-04）
    scores, S_raw, S = compute_pool_scores(panels)
    H, s_series, _s_now = compute_H(panels, target_code, weeks=weeks)
    S_now = float(S[target_code])
    rank = int(S.rank(ascending=False)[target_code])
    sm = side_means(scores.loc[target_code])

    # 2) 复核日（AC-006）与提前复核提示（REQ-003）
    rd, crossed = review_date(score_date, n=5)
    flags = early_review_flags(target, target_code)

    # 3) 解读（T-06，硬门禁内置）
    week_change = (float(target["close"].iloc[-1] / target["close"].iloc[-6] - 1)
                   if len(target) >= 6 else None)
    ctx = DiagnosisContext(
        code=target_code, name=target_name, pool=pool_name,
        observe_date=observe_date.strftime("%Y-%m-%d"),
        S=S_now, H=H, rank_in_pool=rank, pool_size=len(panels),
        side_scores=sm, factor_scores={f: float(scores.loc[target_code][f]) for f in scores.columns},
        factor_labels={f[0]: f[3] for f in FACTOR_TABLE},
        top_positives=[f[3] for f in FACTOR_TABLE
                       if scores.loc[target_code].get(f[0]) is not None
                       and scores.loc[target_code][f[0]] == scores.loc[target_code].max()],
        top_negatives=[f[3] for f in FACTOR_TABLE
                       if scores.loc[target_code].get(f[0]) is not None
                       and scores.loc[target_code][f[0]] <= 35],
        week_change=week_change, review_date=rd.isoformat(),
    )
    interp = get_interpretation(ctx)

    # 4) 图表 4 张
    labels = {f[0]: f[3] for f in FACTOR_TABLE}
    charts = {
        "price": _chart_price(target, target_name),
        "factors": _chart_factors(scores.loc[target_code], labels),
        "series": _chart_s_series(s_series, target_name),
        "hist": _chart_hist(s_series, S_now),
    }

    # 5) 章节内容组装
    interp_paras = interp.text.split("\n")
    side_rows = "".join(
        f"<tr><td>{s}</td><td>{_fmt(sm.get(s), 0)}</td><td style='color:{band(sm.get(s))[1]}'>"
        f"{band(sm.get(s))[2]}</td></tr>" for s in SIDE_ORDER)
    factor_rows = "".join(
        f"<tr><td>{f[0]}</td><td>{f[1]}</td><td>{f[3]}</td><td>{_fmt(scores.loc[target_code][f[0]], 1)}</td>"
        f"<td style='color:{band(scores.loc[target_code][f[0]])[1]}'>{band(scores.loc[target_code][f[0]])[2]}</td></tr>"
        for f in FACTOR_TABLE)
    corr = scores.corr()
    corr_vals = corr.where(~np.eye(len(corr), dtype=bool), np.nan).stack().dropna()
    corr_max = float(corr_vals.max()); corr_mean = float(corr_vals.abs().mean())
    below40_weeks = int((s_series.iloc[:-1] < 40).sum())
    data_start = target.index[0].strftime("%Y-%m-%d")
    n_days = len(target)

    html = f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">
<title>{target_name}（{target_code}）周频因子诊断报告 · 合并版</title>
<style>
body{{font-family:"Microsoft YaHei",sans-serif;max-width:900px;margin:24px auto;padding:0 16px;color:#222;line-height:1.6}}
h1{{font-size:20px;border-bottom:2px solid #333;padding-bottom:8px}}
h2{{font-size:15px;border-left:4px solid #555;padding-left:8px;margin-top:28px}}
table{{border-collapse:collapse;width:100%;font-size:13px}}
td,th{{border:1px solid #ccc;padding:5px 8px}}
.small{{color:#777;font-size:12px}}
.part{{background:#f6f6f6;padding:4px 12px;font-weight:bold;margin-top:34px;border-radius:4px}}
</style></head><body>
<h1>{target_name}（{target_code}）周频因子诊断报告 · 合并版</h1>
<p class="small">生成时间 {_dt.now().strftime("%Y-%m-%d %H:%M")} · 因子清单见 spec v1.3 附表 A（18 项，PM 审定冻结）· 解读生成器：{interp.generator}</p>

<div class="part">简要报告</div>
<h2>{SECTION_TITLES["meta"]}</h2>
<table><tr><th>代码</th><td>{target_code}</td><th>名称</th><td>{target_name}</td></tr>
<tr><th>观察日（出分日）</th><td>{observe_date.strftime("%Y-%m-%d")}</td><th>同池</th><td>{pool_name}（{len(panels)} 只有效样本）</td></tr>
<tr><th>建议复核日</th><td>{rd.isoformat()}（出分日 + 5 个交易日{ '，跨节假日' if crossed else ''}）</td><th>数据级别</th><td>日线级行情与估值（NG-006）</td></tr></table>
<h2>{SECTION_TITLES["summary"]}</h2>
<table><tr><th>综合分 S（同池两级排名）</th><th>历史位置参考分 H</th><th>同池排名</th><th>近一周涨跌幅</th></tr>
<tr><td style="color:{band(S_now)[1]}">{_fmt(S_now, 1)}（{band(S_now)[2]}）</td>
<td style="color:{band(H)[1]}">{_fmt(H, 1)}</td><td>{rank}/{len(panels)}</td><td>{_pct(week_change)}</td></tr></table>
<h2>{SECTION_TITLES["oneline"]}</h2>
<p>{interp_paras[0]}</p>
<h2>{SECTION_TITLES["sides"]}</h2>
<table><tr><th>观察侧面</th><th>侧面均值</th><th>分档</th></tr>{side_rows}</table>
<p class="small">四个观察侧面仅用于拆解展示，不另设权重（BR-006）。</p>
<h2>{SECTION_TITLES["curve"]}</h2><img src="data:image/png;base64,{charts['price']}"/>
<h2>{SECTION_TITLES["howto"]}</h2><p>{HOWTO_TEXT}</p>
<h2>{SECTION_TITLES["plain"]}</h2>{''.join(f'<p>{p}</p>' for p in interp_paras[1:])}
<h2>{SECTION_TITLES["boundary"]}</h2><p>{BOUNDARY_TEXT}</p>
<p class="small">{' '.join(flags)}</p>

<div class="part">详细报告</div>
<h2>{SECTION_TITLES["objective"]}</h2>
<p>截至 {observe_date.strftime("%Y-%m-%d")}，{target_name}（{target_code}）的 18 项日线因子在 {pool_name}
同池（{len(panels)} 只）中的综合分 S={_fmt(S_now, 1)}（第 {rank}/{len(panels)} 名）；其中
{sum(1 for f in FACTOR_TABLE if scores.loc[target_code][f[0]] > 60)} 项处于偏有利区（>60），
{sum(1 for f in FACTOR_TABLE if scores.loc[target_code][f[0]] < 40)} 项处于偏不利区（<40）。
因子原始值基于前复权日线与截面估值计算，缺失项不填分、不计权（BR-006）。</p>
<h2>{SECTION_TITLES["compare"]}</h2>
<p>S 与 H 是两种口径：S={_fmt(S_now, 1)} 为<b>横截面</b>位置（当前因子状态在同池中的相对强弱）；
H={_fmt(H, 1)} 为<b>时间序列</b>位置（当前 S 相对此前 52 个完成周 S 的百分位）。
S 低而 H 高，说明当前虽处同池偏弱位置、但高于自身一年来的多数时点；反之亦然。两者不可相互替代。</p>
<h2>{SECTION_TITLES["dist"]}</h2><img src="data:image/png;base64,{charts['factors']}"/>
<h2>{SECTION_TITLES["dims"]}</h2>
<p>宽口径维度观察：估值特征 3 项（均值 {_fmt(sm['估值特征'], 0)}）、价格行为 7 项（{_fmt(sm['价格行为'], 0)}）、
交易行为 4 项（{_fmt(sm['交易行为'], 0)}）、风险特征 4 项（{_fmt(sm['风险特征'], 0)}），共 18 项（附表 A）。</p>
<h2>{SECTION_TITLES["histpos"]}</h2><img src="data:image/png;base64,{charts['series']}"/>
<p>52 完成周轨迹：历史 S 最小 {s_series.min():.1f} / 中位 {s_series.median():.1f} / 最大 {s_series.max():.1f}；
此前 52 周中低于 40 的周数为 {below40_weeks} 周。</p>
<h2>{SECTION_TITLES["scorepos"]}</h2><img src="data:image/png;base64,{charts['hist']}"/>
<p>当前 S={_fmt(S_now, 1)} 位于历史分布的 {_fmt(H, 1)} 百分位（即 H）。</p>
<h2>{SECTION_TITLES["stat"]}</h2>
<p>同池历史关联：因子分矩阵（{len(panels)}×18）的平均绝对相关系数 {corr_mean:.2f}，最高单对相关 {corr_max:.2f}——
因子间存在一定共线，等权合成未做正交化处理，属口径内的已知简化。解释边界：分数描述因子状态，
不代表上涨概率、可能性或收益预期（BR-003）。</p>
<h2>{SECTION_TITLES["factors"]}</h2>
<table><tr><th>编号</th><th>侧面</th><th>观察项（匿名标签）</th><th>同池百分位</th><th>分档</th></tr>{factor_rows}</table>
<p class="small">因子对外仅展示匿名标签与数学原理示意；内部实现依据 spec 附表 A（2026-10-17 PM 审定冻结）。</p>
<h2>{SECTION_TITLES["math"]}</h2>
<p>数学方法：因子分 = 统一方向后的平均并列名次在同池的百分位（0–100）；综合分 S = 有效因子等权平均后放回同池重新排名；
H = 当前 S 相对此前 52 个完成周 S 的百分位位置。数据覆盖：{data_start} 至 {observe_date.strftime("%Y-%m-%d")}
共 {n_days} 个交易日（日线级，前复权）；估值（市盈率 TTM/市净率/总市值）取截面最新值。</p>
<p class="small">本报告由量化管线自动生成 · {BOUNDARY_TEXT}</p>
</body></html>"""

    # 6) 章节完整性（AC-001）+ 合规硬门禁（AC-005，FAIL 抛异常不落盘）
    missing = check_sections(html)
    if missing:
        raise RuntimeError(f"章节完整性检查失败，缺章: {missing}")
    assert_compliant_or_raise(_html_to_text(html), source="报告生成器")

    html_path = os.path.join(out_dir, f"{target_name}_{target_code}_因子诊断报告_合并版.html")
    docx_path = os.path.join(out_dir, f"{target_name}_{target_code}_因子诊断报告_合并版.docx")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    # Gate V 缺陷修复（2026-10-17）：docx 补齐观察日/四侧面表/18 因子明细表（AC-011 内容一致）
    _render_docx(html_path, docx_path, target_name, target_code, pool_name,
                 S_now, H, rank, len(panels), interp_paras, flags, rd,
                 observe_date=score_date, sides=sm,
                 factor_scores=scores.loc[target_code])
    return {"html": html_path, "docx": docx_path, "S": S_now, "H": H,
            "rank": rank, "pool_size": len(panels), "review_date": rd.isoformat(),
            "interpretation_generator": interp.generator,
            "sections_all_present": True, "compliance": "PASS"}


HTML_SECTION_KEYS = ["meta", "summary", "oneline", "sides", "curve", "howto", "plain",
                     "boundary", "objective", "compare", "dist", "dims", "histpos",
                     "scorepos", "stat", "factors", "math"]


def check_sections(html: str) -> List[str]:
    """AC-001：REQ-002 章节清单一一对应，返回缺失章节标题列表。"""
    return [k for k in HTML_SECTION_KEYS if SECTION_TITLES[k] not in html]


def _html_to_text(html: str) -> str:
    import re
    t = re.sub(r"<style[\s\S]*?</style>", "", html)
    t = re.sub(r"<[^>]+>", "", t)
    return t


def _render_docx(html_path: str, docx_path: str, name: str, code: str, pool: str,
                 S: float, H: float, rank: int, pool_size: int,
                 interp_paras: List[str], flags: List[str], rd,
                 observe_date=None, sides: Optional[Dict] = None,
                 factor_scores=None) -> None:
    """docx 文件版（与 HTML 同源同构；AC-011 内容一致性基础）。

    Gate V 修复（2026-10-17）：补齐观察日、四侧面表、18 项匿名因子明细表，
    使文件版与在线版内容一致（此前三处委托在线版，AC-011 判 FAIL）。
    """
    doc = Document()

    def add_h(t, lvl=2):
        doc.add_heading(t, level=lvl)

    def add_p(t, bold=False, color=None, size=None):
        p = doc.add_paragraph()
        r = p.add_run(t)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:                        # docx 色值不带 # 前缀（SKILL 备忘）
            r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        return p

    def add_table(rows: List[List[str]], header: bool = True):
        tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
        try:
            tb.style = "Table Grid"
        except Exception:
            pass
        for i, rvals in enumerate(rows):
            for j, v in enumerate(rvals):
                cell = tb.rows[i].cells[j]
                cell.text = str(v)
                if header and i == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
        return tb

    # 图表从 HTML 中提取 base64（保证双形态同源）
    import re as _re
    imgs = _re.findall(r'base64,([A-Za-z0-9+/=]+)', open(html_path, encoding="utf-8").read())
    img_bufs = []
    for b in imgs:
        try:
            img_bufs.append(io.BytesIO(base64.b64decode(b)))
        except Exception:
            pass

    add_h(f"{name}（{code}）周频因子诊断报告 · 合并版", 1)
    add_p("简要报告", bold=True)
    add_h(SECTION_TITLES["meta"])
    obs_text = f"观察日（出分日）：{observe_date}" if observe_date is not None \
        else "观察日（出分日）：见在线版元信息"
    add_p(f"代码：{code}　名称：{name}　{obs_text}　同池：{pool}（{pool_size} 只有效样本）"
          f"　建议复核日：{rd.isoformat()}")
    add_h(SECTION_TITLES["summary"])
    sc = band(S); hc = band(H)
    p = doc.add_paragraph()
    r1 = p.add_run(f"综合分 S={S:.1f}（{sc[2]}）")
    r1.font.color.rgb = RGBColor.from_string(sc[1].lstrip("#"))
    r2 = p.add_run(f"　历史位置 H={H:.1f}　同池排名 {rank}/{pool_size}")
    add_h(SECTION_TITLES["oneline"])
    add_p(interp_paras[0] if interp_paras else "")
    add_h(SECTION_TITLES["sides"])
    if sides:
        rows = [["观察侧面", "侧面均值", "分档"]]
        for s_ in SIDE_ORDER:
            v = sides.get(s_)
            rows.append([s_, _fmt(v, 0), band(v)[2]])
        add_table(rows)
        add_p("四个观察侧面仅用于拆解展示，不另设权重（BR-006）。", size=9)
    else:
        add_p("四个观察侧面得分见在线版表格（侧面仅拆解，不另设权重，BR-006）。")
    if img_bufs:
        doc.add_picture(img_bufs[0], width=Pt(430))            # 股性曲线
    add_h(SECTION_TITLES["howto"])
    add_p(HOWTO_TEXT)
    add_h(SECTION_TITLES["plain"])
    for t in (interp_paras[1:] if len(interp_paras) > 1 else []):
        add_p(t)
    add_h(SECTION_TITLES["boundary"])
    add_p(BOUNDARY_TEXT)
    for t in flags:
        add_p(t, size=9)
    doc.add_page_break()
    add_p("详细报告", bold=True)
    add_h(SECTION_TITLES["objective"])
    add_p(f"综合分 S={S:.1f}（{pool} 第 {rank}/{pool_size} 名），18 项日线因子，等权合成两级排名（BR-006 口径）。")
    add_h(SECTION_TITLES["compare"])
    add_p(f"S 为横截面位置（{S:.1f}），H 为时间序列位置（{H:.1f}）；两者口径不同，不可相加或平均。")
    add_h(SECTION_TITLES["dist"])
    if len(img_bufs) > 1:
        doc.add_picture(img_bufs[1], width=Pt(430))            # 因子分布
    add_h(SECTION_TITLES["dims"])
    add_p("宽口径维度观察：估值 3 / 价格 7 / 交易 4 / 风险 4，共 18 项（附表 A）。")
    add_h(SECTION_TITLES["histpos"])
    if len(img_bufs) > 2:
        doc.add_picture(img_bufs[2], width=Pt(430))            # S 轨迹
    add_h(SECTION_TITLES["scorepos"])
    if len(img_bufs) > 3:
        doc.add_picture(img_bufs[3], width=Pt(430))            # H 位置
    add_h(SECTION_TITLES["stat"])
    add_p("因子分矩阵统计与解释边界详见在线版；分数不代表上涨概率、可能性或收益预期（BR-003）。")
    add_h(SECTION_TITLES["factors"])
    if factor_scores is not None:
        rows = [["编号", "侧面", "观察项（匿名标签）", "同池百分位", "分档"]]
        for fid, side, _key, label, _low in FACTOR_TABLE:
            v = float(factor_scores[fid])
            rows.append([fid, side, label, _fmt(v, 1), band(v)[2]])
        add_table(rows)
        add_p("因子对外仅展示匿名标签与数学原理示意；内部实现依据 spec 附表 A（PM 审定冻结）。", size=9)
    else:
        add_p("匿名因子明细 18 项（3 表分组）数值见在线版明细表；对外仅展示匿名标签。")
    add_h(SECTION_TITLES["math"])
    add_p("数学方法：平均并列名次百分位（0–100）→ 等权合成 → 同池两级排名；H 为 52 完成周历史百分位。"
          "数据覆盖：日线级行情与截面估值（NG-006，CR-S-002）。")
    add_p(BOUNDARY_TEXT, size=8)
    doc.save(docx_path)
